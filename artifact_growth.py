from collections import defaultdict, Counter
from itertools import product, permutations, accumulate
from math import prod
from pathlib import Path
from functools import cache
from typing import NamedTuple

import numpy as np
from pandas import DataFrame
from openpyxl import load_workbook
from matplotlib import pyplot as plt


plt.rcParams['font.sans-serif'] = ['SimHei']
growths = {
    '暴击': [2.7, 3.1, 3.5, 3.9],
    '暴伤': [5.4, 6.2, 7.0, 7.8],
    '大攻击': [4.1, 4.7, 5.3, 5.8],
    '小攻击': [14, 16, 18, 19],
    '精通': [16, 19, 21, 23],
    '充能': [4.5, 5.2, 5.8, 6.5],
    '大生命': [4.1, 4.7, 5.3, 5.8],
    '小生命': [209, 239, 269, 299],
    '大防御': [5.1, 5.8, 6.6, 7.3],
    '小防御': [16, 19, 21, 23],
    '': [0, 0, 0, 0],
}


class Item(NamedTuple):
    name: str
    weight: int | float
    init: float | None


@cache
def get_item_dist(init_items: int, init_rand: bool = True) -> dict[tuple, int]:
    if init_items == 3:     # 增加 4 次词条
        dist_freq = {
            (4, 0, 0, 0): 1,
            (3, 1, 0, 0): 4,
            (2, 2, 0, 0): 6,
            (2, 1, 1, 0): 12,
            (1, 1, 1, 1): 24,
        }
    elif init_items == 4:
        dist_freq = {
            (5, 0, 0, 0): 1,
            (4, 1, 0, 0): 5,
            (3, 2, 0, 0): 10,
            (3, 1, 1, 0): 20,
            (2, 2, 1, 0): 30,
            (2, 1, 1, 1): 60,
        }
    else:
        raise ValueError(f'init_items should be 3 or 4, but got {init_items}')

    if init_rand:
        dist_freq = {tuple(i + 1 for i in dist): freq for dist, freq in dist_freq.items()}
    return {d: freq for dist, freq in dist_freq.items() for d in permutations(dist)}


@cache
def get_multi_growth(name: str) -> list:
    return [list(Counter(sum(seq) for seq in product(growths[name], repeat=i)).items()) for i in range(7)]


def calc_score_dist(artifact: list[Item], init_items: int, init_rand: bool = True
                    ) -> tuple[int, list[tuple[float, list[int]]]]:
    item_dist = get_item_dist(init_items, init_rand)
    init_score = 0 if init_rand else sum(item.init * item.weight for item in artifact)
    init_times = sum(bool(item.weight) for item in artifact) if init_rand else 0
    score_dist = defaultdict(lambda: [0] * (2 + init_items))
    multi_growths = {item.name: [[(g * item.weight, t) for g, t in line]
                                 for line in get_multi_growth(item.name)] for item in artifact}
    for dist, freq in item_dist.items():
        growth_seqs = [multi_growths[item.name][d] for d, item in zip(dist, artifact)]
        hits = sum(d for d, item in zip(dist, artifact) if item.weight)
        for growth_seq in product(*growth_seqs):
            score = init_score + sum(g for g, _ in growth_seq)
            score_dist[f'{score:.1f}'][hits - init_times] += freq * prod(t for _, t in growth_seq)
    total = 4 ** (2 + 2 * init_items + 4 * init_rand)
    assert sum(sum(i) for i in score_dist.values()) == total
    return total, sorted(score_dist.items(), key=lambda x: float(x[0]))


def save2excel(name: str, total: int, score_dist,
               as_prob: bool = True, as_str: bool = False) -> None:
    def disp(num: int) -> str | int | float:
        if as_str:
            if as_prob:
                return f'{num/total:.2e}' if num else '0'
            return str(num)
        return num / total if as_prob else num

    path = f'{name}/{name}.xlsx'
    Path(path).parent.mkdir(exist_ok=True, parents=True)
    cols = [score for score, _ in score_dist]
    index = range(len(score_dist[0][1]) - 1, -1, -1)
    lines = [[disp(freq[i]) for _, freq in score_dist] for i in index]
    DataFrame(lines, columns=cols, index=index).to_excel(path)

    wb = load_workbook(path)
    wb.active.cell(row=1, column=1).value = r'命中\得分'
    wb.save(path)


def plot_bar(name: str, total: int, score_dist) -> None:
    path = f'{name}/bar.png'
    Path(path).parent.mkdir(exist_ok=True, parents=True)
    scores = [score for score, _ in score_dist]
    dists = np.array([[dist[i] / total for _, dist in score_dist]
                      for i in range(len(score_dist[0][1]))])
    bottoms = np.zeros(len(scores))
    fig, ax = plt.subplots(figsize=(18, 6))
    for i, dist in enumerate(dists):
        ax.bar(scores, dist, label=f'命中 {i} 次', bottom=bottoms)
        # ax.bar_label(p, label_type='center')
        # 这里 p 是 ax.bar 返回值，但看起来乱，所以注释了
        bottoms += dist
    ax.set_xlabel('得分')
    ax.set_ylabel('概率')
    ax.set_title(name)
    ax.legend()
    ax.set_xticks(range(len(scores)))
    ax.set_xticklabels(scores, rotation=90)
    ax.grid(axis='y')
    plt.savefig(path)
    plt.show()
    plt.close()


def plot_pmf(name: str, total: int, score_dist) -> None:
    path = f'{name}/pmf.png'
    Path(path).parent.mkdir(exist_ok=True, parents=True)
    scores = [float(score) for score, _ in score_dist]
    dists = [sum(dist) / total for _, dist in score_dist]
    plt.plot(scores, dists)
    plt.title(f'{name}-PMF')
    plt.xlim(scores[0], scores[-1])
    plt.ylim(min(dists), max(dists))
    plt.xlabel('得分')
    plt.ylabel('概率')
    plt.grid(True)
    plt.savefig(path)
    plt.show()
    plt.close()


def plot_cdf(name: str, total: int, score_dist) -> None:
    path = f'{name}/cdf.png'
    Path(path).parent.mkdir(exist_ok=True, parents=True)
    scores = [float(score) for score, _ in score_dist]
    dists = list(accumulate([sum(dist) / total for _, dist in score_dist]))

    plt.plot(scores, dists)
    y_alpha = [(0.1, 0.9), (0.3, 0.99), (0.5, 0.999), (0.7, 0.9999), (0.9, 0.99999),
               (0.2, 0.5), (0.15, 0.1)]
    for y, alpha in y_alpha:
        idx = next(i for i, prob in reversed(list(enumerate(dists))) if prob < alpha)
        line, = plt.plot([scores[idx], scores[idx]], [0, dists[idx]], '--')
        color = line.get_color()
        plt.text(scores[idx], y, f'α={alpha}', fontsize=12, ha='right', color=color)
        plt.text(scores[idx], y - 0.05, f'x={scores[idx]}', fontsize=12, ha='right', color=color)
    plt.title(f'{name}-CDF')
    plt.xlim(scores[0], scores[-1])
    plt.ylim(0, 1)
    plt.xlabel('得分')
    plt.ylabel('概率')
    plt.grid(axis='y')
    plt.savefig(path)
    plt.show()
    plt.close()


def get_excel_png(name: str, artifact: list[Item], init_items: int = 4, init_rand: bool = True) -> None:
    data = calc_score_dist(artifact, init_items, init_rand)
    save2excel(name, *data)
    plot_bar(name, *data)
    plot_pmf(name, *data)
    plot_cdf(name, *data)


def main():
    """同时含有百分比和数值的，或有更多项的，得分结果非常多，导致柱状图看起来混乱，此时 CDF 视觉效果较好"""
    # get_excel_png('4r-暴击-爆伤', [
    #     # (名称, 权重, 初始值)
    #     Item(name='暴击', weight=2, init=None),
    #     Item(name='暴伤', weight=1, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=4, init_rand=True)

    # get_excel_png('4r-暴击', [
    #     # (名称, 权重, 初始值)
    #     Item(name='暴击', weight=2, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=4, init_rand=True)

    # get_excel_png('3r-暴击-爆伤', [
    #     # (名称, 权重, 初始值)
    #     Item(name='暴击', weight=2, init=None),
    #     Item(name='暴伤', weight=1, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=3, init_rand=True)

    # get_excel_png('3r-暴击', [
    #     # (名称, 权重, 初始值)
    #     Item(name='暴击', weight=2, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=3, init_rand=True)

    # get_excel_png('4r-大生命-小生命', [
    #     # (名称, 权重, 初始值)
    #     Item(name='大生命', weight=1, init=None),
    #     Item(name='小生命', weight=0.01, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=4, init_rand=True)

    # get_excel_png('4r-大攻击-小攻击', [
    #     # (名称, 权重, 初始值)
    #     Item(name='大攻击', weight=1, init=None),
    #     Item(name='小攻击', weight=0.1, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=4, init_rand=True)

    # get_excel_png('4r-小攻击', [
    #     # (名称, 权重, 初始值)
    #     Item(name='小攻击', weight=1, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    #     Item(name='', weight=0, init=None),
    # ], init_items=4, init_rand=True)


if __name__ == '__main__':
    from timeit import timeit
    print(timeit(main, number=1))
